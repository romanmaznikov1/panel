# -*- coding: utf-8 -*-
"""Работа с исходным .xlsx: только чтение (для разовой миграции в базу)
и запись отчёта-выгрузки (кнопка «Экспорт в Excel»). Сам файл базы данных
в этом модуле не используется — это чисто мост к Excel.
"""

import os
import re
from copy import copy
from datetime import datetime

APP_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(APP_DIR)

import sys  # noqa: E402
sys.path.insert(0, os.path.join(APP_DIR, "vendor"))

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.styles.cell_style import StyleArray  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

EXPORT_FILENAME = "Экспорт для Excel.xlsx"
EXPORT_PATH = os.path.join(ROOT_DIR, EXPORT_FILENAME)

STUDENT_COLUMNS = [
    ("name", 2),
    ("phone", 3),
    ("direction", 4),
    ("start", 5),
    ("plan", 6),
    ("price", 7),
    ("until", 8),
    ("status", 9),
    ("source", 10),
    ("note", 11),
]
NUMERIC_STUDENT_FIELDS = {"price"}
STATUSES = ["Активен", "Продлил", "Ушёл"]

FINANCE_COLUMNS = [
    ("income_subs", 2),
    ("income_rent", 3),
    ("expense_rent", 5),
    ("expense_coaches", 6),
    ("expense_admins", 7),
    ("expense_ads", 8),
    ("expense_merch", 9),
    ("expense_other", 10),
]

MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]

MIN_STUDENT_ROWS = 60
SPARE_STUDENT_ROWS = 10


def source_workbook_path():
    """Находит исходный xlsx-файл (не выгрузку) рядом с папкой app."""
    for name in sorted(os.listdir(ROOT_DIR)):
        if not name.endswith(".xlsx") or name.startswith("~$"):
            continue
        if name == EXPORT_FILENAME:
            continue
        return os.path.join(ROOT_DIR, name)
    return None


# --- чтение (только для миграции) ------------------------------------------

def branch_list(wb):
    branches = []
    for title in wb.sheetnames:
        if title.startswith("Ученики_"):
            slug = title.split("_", 1)[1]
            finance = "Финансы_" + slug
            if finance in wb.sheetnames:
                branches.append({"id": slug, "name": slug, "students": title, "finance": finance})
    return branches


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def cell_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(re.sub(r"[\s  ]", "", str(value)).replace(",", "."))
    except ValueError:
        return None


def is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def parse_ru_date(text):
    """«25.07.2026» -> «2026-07-25». Возвращает "" если не распознано."""
    match = re.match(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$", str(text or "").strip())
    if not match:
        return ""
    day, month, year = (int(x) for x in match.groups())
    try:
        return datetime(year, month, day).strftime("%Y-%m-%d")
    except ValueError:
        return ""


def format_ru_date(iso):
    if not iso:
        return ""
    try:
        return datetime.strptime(iso, "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return ""


def read_students(ws):
    students = []
    for row in range(3, ws.max_row + 1):
        name = ws.cell(row=row, column=2).value
        if is_formula(name):
            continue
        name = cell_text(name)
        if not name:
            continue
        if name.startswith("Сводка") or name.startswith("Всего"):
            break
        item = {}
        for field, col in STUDENT_COLUMNS:
            raw = ws.cell(row=row, column=col).value
            if is_formula(raw):
                raw = None
            if field in NUMERIC_STUDENT_FIELDS:
                item[field] = cell_number(raw)
            else:
                item[field] = cell_text(raw)
        if item["status"] not in STATUSES:
            item["status"] = STATUSES[0]
        students.append(item)
    return students


def read_finance(ws):
    months = []
    for row in range(2, ws.max_row + 1):
        label = cell_text(ws.cell(row=row, column=1).value)
        if not label or label.upper().startswith("ИТОГО"):
            break
        item = {"label": label}
        for field, col in FINANCE_COLUMNS:
            raw = ws.cell(row=row, column=col).value
            item[field] = cell_number(raw) if not is_formula(raw) else None
        months.append(item)
    return months


def read_source_for_migration():
    """Читает исходный xlsx целиком в простую структуру для миграции в БД."""
    path = source_workbook_path()
    if not path:
        return None
    wb = openpyxl.load_workbook(path)
    branches = []
    for meta in branch_list(wb):
        branches.append(
            {
                "id": meta["id"],
                "name": meta["name"],
                "students": read_students(wb[meta["students"]]),
                "months": read_finance(wb[meta["finance"]]),
            }
        )
    wb.close()
    return {"source_path": path, "branches": branches}


def month_to_year_month(label):
    match = re.match(r"([А-Яа-яЁё]+)\s+(\d{4})", label or "")
    if not match:
        return None
    name, year = match.group(1).capitalize(), int(match.group(2))
    if name not in MONTHS_RU:
        return None
    return year, MONTHS_RU.index(name) + 1


# --- запись (только для экспорта-выгрузки) ---------------------------------

def row_style(ws, row, columns):
    return {col: copy(ws.cell(row=row, column=col)._style) for col in columns}


def apply_style(ws, row, styles):
    for col, style in styles.items():
        ws.cell(row=row, column=col)._style = copy(style)


def write_students(ws, students):
    columns = list(range(1, 12))
    data_style = row_style(ws, 4, columns)

    note = ""
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= 3:
            note = note or cell_text(ws.cell(row=merged.min_row, column=merged.min_col).value)
            ws.unmerge_cells(str(merged))

    capacity = max(MIN_STUDENT_ROWS, len(students) + SPARE_STUDENT_ROWS)
    first, last = 3, 2 + capacity

    for row in range(3, max(ws.max_row, last + 12) + 1):
        for col in columns:
            ws.cell(row=row, column=col).value = None

    for idx in range(capacity):
        row = first + idx
        apply_style(ws, row, data_style)
        ws.cell(row=row, column=1).value = f'=IF(B{row}="","",ROW()-2)'
        if idx < len(students):
            student = students[idx]
            for field, col in STUDENT_COLUMNS:
                value = student.get(field)
                if field in NUMERIC_STUDENT_FIELDS:
                    ws.cell(row=row, column=col).value = value if value not in (None, "") else None
                elif field in ("start", "until"):
                    ws.cell(row=row, column=col).value = format_ru_date(value) or None
                else:
                    ws.cell(row=row, column=col).value = value or None

    for row in range(last + 1, max(ws.max_row, last + 12) + 1):
        for col in columns:
            ws.cell(row=row, column=col)._style = StyleArray()

    stats = last + 3
    ws.cell(row=stats, column=2).value = "Сводка по филиалу"
    ws.cell(row=stats, column=2).font = Font(bold=True, color="FF1F4E78")
    rows = [
        ("Всего учеников (заполнено):", f"=COUNTA(B{first}:B{last})"),
        ("Активны:", f'=COUNTIF(I{first}:I{last},"Активен")'),
        ("Продлили:", f'=COUNTIF(I{first}:I{last},"Продлил")'),
        ("Ушли:", f'=COUNTIF(I{first}:I{last},"Ушёл")'),
        (
            "% удержания (продлили / (продлили+ушли)):",
            f'=IFERROR(COUNTIF(I{first}:I{last},"Продлил")/'
            f'(COUNTIF(I{first}:I{last},"Продлил")+COUNTIF(I{first}:I{last},"Ушёл")),0)',
        ),
    ]
    for offset, (label, formula) in enumerate(rows, start=1):
        ws.cell(row=stats + offset, column=2).value = label
        cell = ws.cell(row=stats + offset, column=3)
        cell.value = formula
        if label.startswith("%"):
            cell.number_format = "0%"

    if note:
        note_row = stats + 7
        ws.cell(row=note_row, column=2).value = note
        ws.cell(row=note_row, column=2).font = Font(italic=True, color="FF7F7F7F")
        ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)

    ws.data_validations.dataValidation = []
    validation = DataValidation(type="list", formula1='"Активен,Продлил,Ушёл"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"I{first}:I{last}")

    return {"first": first, "last": last, "stats": stats + 1}


def write_finance(ws, months):
    columns = list(range(1, 13))
    data_style = row_style(ws, 2, columns)
    total_style = row_style(ws, 15, columns) if ws.max_row >= 15 else data_style

    for row in range(2, ws.max_row + 6):
        for col in columns:
            ws.cell(row=row, column=col).value = None

    for idx, month in enumerate(months):
        row = 2 + idx
        apply_style(ws, row, data_style)
        ws.cell(row=row, column=1).value = month.get("label") or ""
        for field, col in FINANCE_COLUMNS:
            value = month.get(field)
            ws.cell(row=row, column=col).value = value if value not in (None, "") else None
        ws.cell(row=row, column=4).value = f"=B{row}+C{row}"
        ws.cell(row=row, column=11).value = f"=SUM(E{row}:J{row})"
        ws.cell(row=row, column=12).value = f"=D{row}-K{row}"

    first, last = 2, 1 + len(months)
    total = last + 2
    apply_style(ws, total, total_style)
    ws.cell(row=total, column=1).value = "ИТОГО за период"
    for col in range(2, 13):
        letter = openpyxl.utils.get_column_letter(col)
        ws.cell(row=total, column=col).value = f"=SUM({letter}{first}:{letter}{last})"
    return {"total": total}


def write_summary(ws, branches, layout):
    for idx, branch in enumerate(branches):
        row = 6 + idx
        students_sheet = f"'{branch['students_sheet']}'"
        finance_sheet = f"'{branch['finance_sheet']}'"
        stats = layout[branch["id"]]["students"]["stats"]
        ws.cell(row=row, column=1).value = branch["name"]
        for offset, col in enumerate(range(2, 7)):
            ws.cell(row=row, column=col).value = f"={students_sheet}!C{stats + offset}"
        total = layout[branch["id"]]["finance"]["total"]
        frow = 13 + idx
        ws.cell(row=frow, column=1).value = branch["name"]
        ws.cell(row=frow, column=2).value = f"={finance_sheet}!D{total}"
        ws.cell(row=frow, column=3).value = f"={finance_sheet}!K{total}"
        ws.cell(row=frow, column=4).value = f"={finance_sheet}!L{total}"


def export_to_xlsx(data):
    """Строит выгрузку-отчёт на основе исходного шаблона и данных из БД.
    Исходный файл открывается только на чтение (как шаблон стилей/формул),
    результат сохраняется отдельным файлом — оригинал не трогается.
    """
    source = source_workbook_path()
    if not source:
        return export_plain_xlsx(data)
    wb = openpyxl.load_workbook(source)
    metas = {m["id"]: m for m in branch_list(wb)}
    layout = {}
    branches_for_summary = []
    for branch in data["branches"]:
        meta = metas.get(branch["id"])
        if not meta:
            continue
        layout[branch["id"]] = {
            "students": write_students(wb[meta["students"]], branch["students"]),
            "finance": write_finance(wb[meta["finance"]], branch["months"]),
        }
        branches_for_summary.append(
            {
                "id": branch["id"],
                "name": branch["name"],
                "students_sheet": meta["students"],
                "finance_sheet": meta["finance"],
            }
        )
    if "Сводка" in wb.sheetnames:
        write_summary(wb["Сводка"], branches_for_summary, layout)
    tmp = EXPORT_PATH + ".tmp"
    wb.save(tmp)
    wb.close()
    os.replace(tmp, EXPORT_PATH)
    return EXPORT_PATH


# --- выгрузка без шаблона ----------------------------------------------------

FINANCE_TITLES = [
    ("income_subs", "Абонементы"), ("income_rent", "Аренда зала"),
    ("expense_rent", "Аренда помещения"), ("expense_coaches", "ЗП тренеров"),
    ("expense_admins", "ЗП админов"), ("expense_ads", "Реклама"),
    ("expense_merch", "Костюмы"), ("expense_other", "Прочее"),
]
INCOME_FIELDS = {"income_subs", "income_rent"}


def export_plain_xlsx(data):
    """Отчёт с нуля, когда исходного шаблона нет: ученики, финансы, оплаты."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="D42A90")
    names = {b["id"]: b["name"] for b in data["branches"]}

    def sheet(ws, headers, rows, money_cols=()):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for row in rows:
            ws.append(row)
        for idx, title in enumerate(headers, start=1):
            letter = get_column_letter(idx)
            width = max([len(str(title))] + [len(str(r[idx - 1] or "")) for r in rows[:300]]) + 2
            ws.column_dimensions[letter].width = min(width, 48)
            if idx in money_cols:
                for cell in ws[letter][1:]:
                    cell.number_format = '# ##0 "₽"'
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()
    students = []
    for branch in data["branches"]:
        for st in branch["students"]:
            students.append([
                st["name"], st["phone"], branch["name"], st["direction"], st["plan"],
                st.get("price") or None, st["start"], st["until"], st.get("remaining"),
                st["status"], st.get("activity"), st.get("lastVisit"), st.get("paidTotal"),
                st["source"], st["note"],
            ])
    sheet(wb.active, ["Ученик", "Телефон", "Филиал", "Направления", "Абонемент", "Своя цена",
                      "Пришёл впервые", "Действует до", "Осталось занятий", "Статус", "Активность",
                      "Последний визит", "Оплатил всего", "Источник", "Заметка"], students, money_cols=(6, 13))
    wb.active.title = "Ученики"

    finance = []
    for branch in data["branches"]:
        for m in branch["months"]:
            income = sum(m.get(k) or 0 for k, _ in FINANCE_TITLES if k in INCOME_FIELDS)
            expense = sum(m.get(k) or 0 for k, _ in FINANCE_TITLES if k not in INCOME_FIELDS)
            finance.append([branch["name"], m["label"]] + [m.get(k) for k, _ in FINANCE_TITLES]
                           + [income, expense, income - expense])
    if finance:
        sheet(wb.create_sheet("Финансы"), ["Филиал", "Месяц"] + [t for _, t in FINANCE_TITLES]
              + ["Доход", "Расход", "Прибыль"], finance, money_cols=range(3, 14))

    payments = [
        [p["payment_date"], p.get("student_name"), names.get(p["branch_id"], p["branch_id"]), p["amount"], p["note"]]
        for p in sorted(data.get("payments", []), key=lambda p: p["payment_date"], reverse=True)
    ]
    sheet(wb.create_sheet("Оплаты"), ["Дата", "Ученик", "Филиал", "Сумма", "Комментарий"], payments, money_cols=(4,))

    tmp = EXPORT_PATH + ".tmp"
    wb.save(tmp)
    wb.close()
    os.replace(tmp, EXPORT_PATH)
    return EXPORT_PATH
